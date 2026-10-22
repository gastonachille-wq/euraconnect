from __future__ import annotations
import base64, json, os, tempfile, threading, webbrowser, socket, time, unicodedata, io, copy, re
from datetime import date, datetime, timedelta
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse, parse_qs
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from supabase import create_client

ROOT = Path(__file__).resolve().parent
CONFIG = json.loads((ROOT / "config.json").read_text(encoding="utf-8"))
SUPABASE_URL = CONFIG["supabase_url"]
SUPABASE_KEY = CONFIG["supabase_publishable_key"]
REFRESH_SECONDS = int(CONFIG.get("refresh_seconds", 15))

STATUSES = [
    "1. Prise de contact", "2. À recontacter", "3. Intéressé 2",
    "4. Rdv réalisé", "5. Intéressé 1", "6. Proposition envoyée",
    "7. Accord de principe", "8. Gagné (fermé)", "9. Perdu (fermé)"
]
ODJ_VALUES = ["", "AL", "LS", "NG"]
ODJ_NAMES = {"AL":"Aurélien", "LS":"Lise", "NG":"Nathalie"}

class AppState:
    def __init__(self):
        self.lock = threading.RLock()
        self.client = create_client(SUPABASE_URL, SUPABASE_KEY)
        self.email = None
        self.imports = {}

    def login(self, email, password):
        with self.lock:
            # recreate client to avoid inheriting a stale session
            self.client = create_client(SUPABASE_URL, SUPABASE_KEY)
            result = self.client.auth.sign_in_with_password({"email": email, "password": password})
            if not getattr(result, "user", None):
                raise RuntimeError("Connexion refusée par Supabase.")
            self.email = result.user.email or email
            return {"email": self.email}

    def logout(self):
        with self.lock:
            try: self.client.auth.sign_out()
            except Exception: pass
            self.client = create_client(SUPABASE_URL, SUPABASE_KEY)
            self.email = None

    def require_auth(self):
        if not self.email:
            raise PermissionError("AUTH_REQUIRED")

    def fetch_all(self, table, order=None):
        self.require_auth()
        rows, start, size = [], 0, 1000
        while True:
            q = self.client.table(table).select("*")
            if order:
                q = q.order(order, desc=False)
            resp = q.range(start, start + size - 1).execute()
            batch = resp.data or []
            rows.extend(batch)
            if len(batch) < size: break
            start += size
        return rows

    def history(self, **kw):
        kw.setdefault("utilisateur", self.email or "")
        self.client.table("historique").insert(kw).execute()

STATE = AppState()


EXPORT_TEMPLATE = ROOT / "export_template.xlsx"


def _as_date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s[:19], fmt).date()
        except Exception:
            pass
    return None


def _copy_row_style(ws, src_row, dst_row, max_col):
    """Copie uniquement la présentation de la ligne modèle, sans les valeurs."""
    for c in range(1, max_col + 1):
        src = ws.cell(src_row, c)
        dst = ws.cell(dst_row, c)
        if src.has_style:
            dst._style = copy.copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy.copy(src.alignment)
        if src.protection:
            dst.protection = copy.copy(src.protection)
    if src_row in ws.row_dimensions:
        ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def _clear_data_rows(ws, max_col, template_row=2):
    """Vide la zone de données tout en conservant l'en-tête et une ligne modèle."""
    last = max(ws.max_row, 2)
    # On garde la mise en forme de la ligne 2 comme modèle, mais on efface toutes les valeurs.
    for r in range(2, last + 1):
        for c in range(1, max_col + 1):
            ws.cell(r, c).value = None
    return template_row


def _resize_table(ws, table_name, last_row, last_col):
    if table_name in ws.tables:
        from openpyxl.utils import get_column_letter
        ws.tables[table_name].ref = f"A1:{get_column_letter(last_col)}{max(last_row,2)}"


def _safe_sheet_title(name, used):
    title = re.sub(r'[\\/*?:\[\]]', ' ', str(name or 'Sourcing')).strip() or 'Sourcing'
    title = title[:31]
    base = title
    i = 2
    while title in used:
        suffix = f" {i}"
        title = (base[:31-len(suffix)] + suffix)
        i += 1
    used.add(title)
    return title


def build_xlsx_export():
    """Construit un XLSX compatible avec le fichier d'import initial."""
    STATE.require_auth()
    if not EXPORT_TEMPLATE.exists():
        raise RuntimeError("Modèle XLSX d'export introuvable.")

    wb = load_workbook(EXPORT_TEMPLATE, data_only=False)
    prospects = STATE.fetch_all('prospects', 'nom_sponsor')
    projects = STATE.fetch_all('projets', 'nom_projet')
    sourcing = STATE.fetch_all('sourcing', 'created_at')

    # ------------------------------------------------------------------
    # Reporting : format exact de l'input initial A:I
    # ------------------------------------------------------------------
    ws = wb['Reporting']
    _clear_data_rows(ws, 9)
    for idx, x in enumerate(prospects, start=2):
        if idx > 2:
            _copy_row_style(ws, 2, idx, 9)
        values = [
            x.get('nom_sponsor'), x.get('statut'), x.get('commentaire'),
            _as_date(x.get('date_maj')), x.get('montant'), x.get('referent'),
            x.get('projet_cible'), x.get('odj'), x.get('lien_crm')
        ]
        for c, v in enumerate(values, start=1):
            ws.cell(idx, c).value = v
        ws.cell(idx, 1).font = copy.copy(ws.cell(idx, 1).font)
        ws.cell(idx, 1).font = ws.cell(idx, 1).font.copy(bold=True)
        ws.cell(idx, 4).number_format = '[$-fr-FR]dddd d mmmm yyyy'
    _resize_table(ws, 'Tableau1', len(prospects) + 1, 9)

    # ------------------------------------------------------------------
    # Projets à sponsoriser : format initial A:M
    # ------------------------------------------------------------------
    ws = wb['Projets à sponsoriser']
    _clear_data_rows(ws, 13)
    for idx, x in enumerate(projects, start=2):
        if idx > 2:
            _copy_row_style(ws, 2, idx, 13)
        raw = x.get('donnees_source') or {}
        project_id = raw.get('ID projet') or raw.get('ID Projet') or (idx - 1)
        values = [
            project_id, x.get('nom_projet'), x.get('priorite'), x.get('odj'),
            x.get('referent'), x.get('montant_cible'), _as_date(x.get('date_projet')),
            x.get('partenaire_cible'), x.get('site'), x.get('campagne_crm'),
            x.get('commentaires'), x.get('slides'), x.get('finance')
        ]
        for c, v in enumerate(values, start=1):
            ws.cell(idx, c).value = v
        if values[6]:
            ws.cell(idx, 7).number_format = 'dd/mm/yyyy'
    _resize_table(ws, 'Tableau2', len(projects) + 1, 13)

    # ------------------------------------------------------------------
    # Sourcing : reconstitution par onglet source, à partir du JSON brut.
    # Les onglets connus conservent leurs en-têtes / styles du modèle.
    # ------------------------------------------------------------------
    grouped = {}
    for x in sourcing:
        source = str(x.get('source') or 'Sourcing')
        source = source.split(' [ligne sans nom]')[0]
        grouped.setdefault(source, []).append(x)

    protected = {
        'Reporting', 'Projets à sponsoriser', 'Ordre du jour (automatique)',
        'ODJ-28.04', 'ODJ-26.03', 'Listes', 'Mapping Projet <> Prospect', 'Feuil1'
    }
    used = set(wb.sheetnames)

    for source, rows in grouped.items():
        if source in protected:
            # Les lignes de sourcing issues d'un onglet métier protégé ne doivent pas écraser cet onglet.
            source_title = _safe_sheet_title(f"Sourcing - {source}", used)
            ws = wb.create_sheet(source_title)
            existing_headers = []
        elif source in wb.sheetnames:
            ws = wb[source]
            existing_headers = [ws.cell(1,c).value for c in range(1, ws.max_column+1)]
            _clear_data_rows(ws, ws.max_column)
        else:
            source_title = _safe_sheet_title(source, used)
            ws = wb.create_sheet(source_title)
            existing_headers = []

        # Union des clés en conservant l'ordre du modèle si disponible.
        keys = []
        seen = set()
        for h in existing_headers:
            if h not in (None, ''):
                keys.append(str(h)); seen.add(str(h))
        for item in rows:
            for k in (item.get('donnees') or {}).keys():
                if str(k) not in seen:
                    keys.append(str(k)); seen.add(str(k))
        if not keys:
            keys = ['Nom']
        for c, h in enumerate(keys, start=1):
            ws.cell(1,c).value = h
        for r_idx, item in enumerate(rows, start=2):
            if r_idx > 2 and ws.max_row >= 2:
                _copy_row_style(ws, 2, r_idx, len(keys))
            data = item.get('donnees') or {}
            for c, h in enumerate(keys, start=1):
                ws.cell(r_idx,c).value = data.get(h)
        # Redimensionner les tables connues.
        if source == 'Santexpo26':
            _resize_table(ws, 'Tableau3', len(rows)+1, len(keys))

    # ------------------------------------------------------------------
    # Rattachements multiprojets : extension de l'export sans casser Reporting A:I
    # ------------------------------------------------------------------
    if 'Rattachements multiprojets' in wb.sheetnames:
        del wb['Rattachements multiprojets']
    wr = wb.create_sheet('Rattachements multiprojets')
    wr.append(['Prospect','Projet','Projet actif'])
    try:
        rels=STATE.fetch_all('prospect_projets')
        p_by_id={str(x.get('id')):x for x in prospects}
        j_by_id={str(x.get('id')):x for x in projects}
        for rel in rels:
            prospect=p_by_id.get(str(rel.get('prospect_id')))
            project=j_by_id.get(str(rel.get('projet_id')))
            if prospect and project:
                wr.append([prospect.get('nom_sponsor'),project.get('nom_projet'),'Oui' if _project_active_value(project) else 'Non'])
    except Exception:
        pass

    # Demander à Excel de recalculer les formules (dont l'ODJ automatique) à l'ouverture.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = 'auto'
    except Exception:
        pass

    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def norm(v):
    if v is None: return ""
    s = str(v).strip().lower()
    s = ''.join(c for c in unicodedata.normalize('NFKD', s) if not unicodedata.combining(c))
    return ' '.join(s.replace('_',' ').replace('-',' ').split())


def iso_date(v):
    if v is None or v == "": return None
    if isinstance(v, datetime): return v.date().isoformat()
    if isinstance(v, date): return v.isoformat()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S"):
        try: return datetime.strptime(s[:19], fmt).date().isoformat()
        except Exception: pass
    return None


def num(v):
    if v is None or v == "": return None
    if isinstance(v, (int,float)): return float(v)
    try: return float(str(v).replace('€','').replace(' ','').replace(',','.'))
    except Exception: return None


def canonical_status(v):
    """Mappe les variantes historiques par numero d'etape sans perdre la valeur brute."""
    if v is None or str(v).strip()=="": return None
    raw=str(v).strip()
    if raw in STATUSES: return raw
    import re
    m=re.match(r"\s*([1-9])(?:[\s\.\-_:)]|$)", raw)
    if m:
        return STATUSES[int(m.group(1))-1]
    return None

def canonical_odj(v):
    if v is None: return None
    raw=str(v).strip().upper()
    return raw if raw in ("AL","LS","NG") else None

def row_payload(headers, row):
    return {str(headers[i] if headers[i] is not None else f"Colonne {i+1}"): value_json(v)
            for i,v in enumerate(row) if v not in (None,"")}

def nonempty_update(new, old=None):
    """En mode fusion, une cellule source vide n'efface jamais une donnee deja presente."""
    out={}
    for k,v in new.items():
        if k in ("donnees_source","source_onglet","source_ligne"):
            out[k]=v
        elif v not in (None, ""):
            out[k]=v
    return out

def value_json(v):
    if isinstance(v, (datetime,date)): return v.isoformat()
    return v


def map_headers(headers):
    return {norm(h): i for i,h in enumerate(headers)}


def pick(row, hm, aliases):
    for a in aliases:
        i = hm.get(norm(a))
        if i is not None and i < len(row): return row[i]
    return None


def inspect_excel(path: Path):
    wb = load_workbook(path, read_only=True, data_only=False)
    sheets=[]
    for ws in wb.worksheets:
        vals = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6), values_only=True))
        headers = [str(x).strip() if x is not None else f"Colonne {i+1}" for i,x in enumerate(vals[0] if vals else [])]
        sample = [[value_json(x) for x in r] for r in vals[1:4]]
        role = "ignore"
        n = norm(ws.title)
        if n == "reporting": role="prospects"
        elif "projets a sponsoriser" in n or "projet a sponsoriser" in n: role="projects"
        elif ws.max_row > 1: role="sourcing"
        sheets.append({"name":ws.title,"rows":max(ws.max_row-1,0),"cols":ws.max_column,"headers":headers,"sample":sample,"suggested_role":role})
    wb.close()
    return sheets


def parse_excel(path: Path, mapping):
    wb = load_workbook(path, read_only=True, data_only=False)
    out={"prospects":[],"projects":[],"sourcing":[],"stats":{}}
    for sheet, role in mapping.items():
        if role == "ignore" or sheet not in wb.sheetnames: continue
        ws=wb[sheet]
        it=ws.iter_rows(values_only=True)
        try: headers=list(next(it))
        except StopIteration: continue
        hm=map_headers(headers)
        stat={"excel_rows":0,"non_empty_rows":0,"mapped":0,"preserved_as_sourcing":0,"invalid_status":0,"invalid_odj":0,"invalid_date":0}
        for excel_row_num,row in enumerate(it,start=2):
            stat["excel_rows"] += 1
            if not any(v not in (None,"") for v in row): continue
            stat["non_empty_rows"] += 1
            raw=row_payload(headers,row)
            if role=="prospects":
                name=pick(row,hm,["Nom sponsor","Sponsor","Entreprise","Société","Societe","Nom"])
                if not name:
                    out["sourcing"].append({"source":f"{sheet} [ligne sans nom]","nom_cible":f"Ligne {excel_row_num}","donnees":raw,"importe":False})
                    stat["preserved_as_sourcing"] += 1
                    continue
                raw_status=pick(row,hm,["Statut"]); st=canonical_status(raw_status)
                if raw_status not in (None,"") and st is None: stat["invalid_status"] += 1
                raw_odj=pick(row,hm,["ODJ"]); odj=canonical_odj(raw_odj)
                if raw_odj not in (None,"") and odj is None: stat["invalid_odj"] += 1
                raw_date=pick(row,hm,["Date MAJ","Date de MAJ","Date mise à jour"]); dt=iso_date(raw_date)
                if raw_date not in (None,"") and dt is None: stat["invalid_date"] += 1
                d={
                    "nom_sponsor":str(name).strip(),
                    "statut":st,
                    "commentaire":pick(row,hm,["Commentaire","Commentaires"]),
                    "date_maj":dt,
                    "montant":num(pick(row,hm,["Montant","Montant cible"])),
                    "referent":pick(row,hm,["Référent","Referent"]),
                    "projet_cible":pick(row,hm,["Projet Cible","Projet cible","Projet"]),
                    "odj":odj,
                    "lien_crm":pick(row,hm,["Lien CRM","CRM"]),
                    "donnees_source":raw,
                    "source_onglet":sheet,
                    "source_ligne":excel_row_num
                }
                out["prospects"].append(d); stat["mapped"] += 1
            elif role=="projects":
                name=pick(row,hm,["Projet","Nom projet","Projet à sponsoriser","Nom"])
                if not name:
                    out["sourcing"].append({"source":f"{sheet} [ligne sans nom]","nom_cible":f"Ligne {excel_row_num}","donnees":raw,"importe":False})
                    stat["preserved_as_sourcing"] += 1
                    continue
                raw_odj=pick(row,hm,["ODJ"]); odj=canonical_odj(raw_odj)
                if raw_odj not in (None,"") and odj is None: stat["invalid_odj"] += 1
                raw_date=pick(row,hm,["Date","Date projet","Date évènement","Date evenement"]); dt=iso_date(raw_date)
                if raw_date not in (None,"") and dt is None: stat["invalid_date"] += 1
                d={
                    "nom_projet":str(name).strip(),
                    "priorite":pick(row,hm,["Priorité","Priorite"]),
                    "referent":pick(row,hm,["Référent","Referent"]),
                    "montant_cible":num(pick(row,hm,["Montant à financer","Montant cible","Montant"])),
                    "date_projet":dt,
                    "partenaire_cible":pick(row,hm,["Partenaire ciblé","Partenaire cible","Partenaires ciblés"]),
                    "site":pick(row,hm,["Site"]),
                    "campagne_crm":pick(row,hm,["Campagne CRM"]),
                    "commentaires":pick(row,hm,["Commentaires","Commentaire"]),
                    "slides":pick(row,hm,["Slides","Lien slides"]),
                    "finance":pick(row,hm,["Financé","Finance","Financement"]),
                    "odj":odj,
                    "donnees_source":raw,
                    "source_onglet":sheet,
                    "source_ligne":excel_row_num
                }
                raw_active=pick(row,hm,["Actif","Active","Projet actif","Etat projet","État projet"])
                if raw_active not in (None, ""):
                    d["actif"] = str(raw_active).strip().lower() not in {"0","false","faux","non","no","inactive","inactif","désactivé","desactive"}
                out["projects"].append(d); stat["mapped"] += 1
            else:
                name=pick(row,hm,["Société","Societe","Entreprise","Nom","Company","Exposant","Organisation"])
                if not name: name=next((v for v in row if isinstance(v,str) and v.strip()), None)
                label=str(name).strip() if name else f"Ligne {excel_row_num}"
                out["sourcing"].append({"source":sheet,"nom_cible":label,"donnees":raw,"importe":False})
                stat["mapped"] += 1
        out["stats"][sheet]=stat
    wb.close()
    return out

def clean_payload(table, data):
    allowed={
        "prospects":{"nom_sponsor","statut","commentaire","date_maj","montant","referent","projet_cible","odj","lien_crm"},
        "projets":{"nom_projet","priorite","referent","montant_cible","date_projet","partenaire_cible","site","campagne_crm","commentaires","slides","finance","odj","actif"}
    }[table]
    d={k:v for k,v in data.items() if k in allowed}
    for k in ("date_maj","date_projet"):
        if k in d: d[k]=iso_date(d[k])
    for k in ("montant","montant_cible"):
        if k in d: d[k]=num(d[k])
    if "odj" in d and d["odj"] not in ODJ_VALUES: d["odj"]=""
    if table=="projets" and "actif" in d:
        d["actif"] = bool(d["actif"]) if isinstance(d["actif"], bool) else str(d["actif"]).strip().lower() in {"1","true","vrai","oui","yes","active","actif"}
    if table=="prospects" and "statut" in d and d["statut"] not in STATUSES: raise ValueError("Statut non autorisé")
    return d

def active_project_by_name(name):
    """Retourne le projet actif correspondant au libellé, ou None."""
    target=norm(name)
    if not target: return None
    for project in STATE.fetch_all('projets'):
        if norm(project.get('nom_projet'))==target and project.get('actif', True) is not False:
            return project
    return None

def validate_new_project_attachment(name):
    if not str(name or '').strip(): return
    if not active_project_by_name(name):
        raise ValueError("Ce projet est désactivé ou introuvable. Seuls les projets actifs peuvent être rattachés.")



def _project_active_value(project):
    if not project:
        return False
    v=project.get('actif', True)
    if v is False or v == 0:
        return False
    if isinstance(v,str) and v.strip().lower() in {'false','0','non','no','inactive','inactif','désactivé','desactive'}:
        return False
    return True


def _project_index(projects=None):
    projects = projects if projects is not None else STATE.fetch_all('projets')
    return {str(p.get('id')): p for p in projects if p.get('id')}


def _relations_for_prospect(prospect_id):
    return STATE.client.table('prospect_projets').select('prospect_id,projet_id,created_at').eq('prospect_id',prospect_id).execute().data or []


def enrich_prospects_with_projects(prospects, projects):
    """Ajoute projets[] / projet_ids[] aux prospects sans casser projet_cible historique."""
    try:
        relations = STATE.fetch_all('prospect_projets')
    except Exception as exc:
        msg=str(exc)
        if 'prospect_projets' in msg or 'PGRST' in msg:
            raise RuntimeError("Migration multiprojets manquante. Exécute MIGRATION_MULTIPROJETS.sql dans Supabase puis relance l'application.") from exc
        raise
    pmap={str(p.get('id')):p for p in projects if p.get('id')}
    by_prospect={}
    for rel in relations:
        pid=str(rel.get('prospect_id') or '')
        jid=str(rel.get('projet_id') or '')
        project=pmap.get(jid)
        if pid and project:
            by_prospect.setdefault(pid,[]).append({
                'id': project.get('id'),
                'nom_projet': project.get('nom_projet'),
                'actif': _project_active_value(project),
                'priorite': project.get('priorite'),
            })
    by_name={norm(p.get('nom_projet')):p for p in projects if p.get('nom_projet')}
    for prospect in prospects:
        linked=by_prospect.get(str(prospect.get('id')),[])
        # Compatibilité : si une ancienne valeur projet_cible n'a pas encore de relation, on l'affiche quand même.
        legacy=prospect.get('projet_cible')
        if legacy and not any(norm(x.get('nom_projet'))==norm(legacy) for x in linked):
            p=by_name.get(norm(legacy))
            linked.append({'id':p.get('id') if p else None,'nom_projet':legacy,'actif':_project_active_value(p) if p else False,'priorite':p.get('priorite') if p else None})
        linked.sort(key=lambda x: norm(x.get('nom_projet')))
        prospect['projets']=linked
        prospect['projet_ids']=[str(x['id']) for x in linked if x.get('id')]
    return prospects


def _normalise_project_ids(ids):
    if ids is None:
        return None
    if not isinstance(ids,list):
        raise ValueError('projet_ids doit être une liste')
    return list(dict.fromkeys(str(x) for x in ids if x))


def _validate_added_project_ids(project_ids, existing_ids=None):
    project_ids=_normalise_project_ids(project_ids) or []
    existing=set(str(x) for x in (existing_ids or []))
    projects=_project_index()
    missing=[]; inactive=[]
    for jid in project_ids:
        if jid in existing:
            continue
        p=projects.get(jid)
        if not p:
            missing.append(jid)
        elif not _project_active_value(p):
            inactive.append(p.get('nom_projet') or jid)
    if missing:
        raise ValueError('Projet introuvable : '+', '.join(missing))
    if inactive:
        raise ValueError('Projet(s) désactivé(s) non rattachable(s) : '+', '.join(inactive))
    return project_ids


def _primary_project_name(project_ids, old_primary=''):
    project_ids=_normalise_project_ids(project_ids) or []
    if not project_ids:
        return None
    projects=_project_index()
    old=norm(old_primary)
    for jid in project_ids:
        p=projects.get(jid)
        if p and norm(p.get('nom_projet'))==old:
            return p.get('nom_projet')
    # Si le projet principal historique n'est plus sélectionné, privilégie un projet actif.
    for jid in project_ids:
        p=projects.get(jid)
        if p and _project_active_value(p):
            return p.get('nom_projet')
    for jid in project_ids:
        p=projects.get(jid)
        if p:
            return p.get('nom_projet')
    return None


def sync_prospect_projects(prospect_id, project_ids, old_primary=''):
    """Synchronise la table many-to-many. Les liens inactifs existants peuvent être conservés ou retirés, mais jamais ajoutés."""
    project_ids=_normalise_project_ids(project_ids)
    if project_ids is None:
        return None
    current=_relations_for_prospect(prospect_id)
    current_ids={str(r.get('projet_id')) for r in current if r.get('projet_id')}
    desired=set(project_ids)
    _validate_added_project_ids(project_ids,current_ids)
    to_remove=current_ids-desired
    to_add=desired-current_ids
    for jid in to_remove:
        STATE.client.table('prospect_projets').delete().eq('prospect_id',prospect_id).eq('projet_id',jid).execute()
    for jid in to_add:
        STATE.client.table('prospect_projets').insert({'prospect_id':prospect_id,'projet_id':jid}).execute()
    primary=_primary_project_name(project_ids,old_primary)
    STATE.client.table('prospects').update({'projet_cible':primary}).eq('id',prospect_id).execute()
    return {'added':list(to_add),'removed':list(to_remove),'primary':primary}


def backfill_legacy_project_links():
    """Crée les relations multiprojets manquantes à partir du champ historique projet_cible."""
    projects=STATE.fetch_all('projets')
    by_name={norm(p.get('nom_projet')):p for p in projects if p.get('nom_projet')}
    existing={(str(r.get('prospect_id')),str(r.get('projet_id'))) for r in STATE.fetch_all('prospect_projets')}
    added=0
    for prospect in STATE.fetch_all('prospects'):
        name=prospect.get('projet_cible')
        project=by_name.get(norm(name)) if name else None
        key=(str(prospect.get('id')),str(project.get('id'))) if project else None
        if project and key not in existing:
            STATE.client.table('prospect_projets').insert({'prospect_id':prospect['id'],'projet_id':project['id']}).execute()
            existing.add(key); added+=1
    return added

def find_free_port(start=8765):
    for p in range(start,start+50):
        with socket.socket() as s:
            try: s.bind(('127.0.0.1',p)); return p
            except OSError: pass
    raise RuntimeError("Aucun port local disponible")

def safe_fetch_all(table, order=None):
    try:
        return STATE.fetch_all(table, order), True
    except Exception as exc:
        msg=str(exc).lower()
        if 'pgrst' in msg or 'does not exist' in msg or 'relation' in msg or 'schema cache' in msg:
            return [], False
        raise

# Structures de l'écosystème Eurasante à exclure du référentiel Relations / RDV.
# Elles restent évidemment intactes dans les autres tables du CRM.
RELATION_EXCLUDED_DOMAINS={
    'eurasante.com','eurasante.onmicrosoft.com',
    'clubster-nsl.com','eurasenior.com','euralimentaire.com','vivalley-campus.com'
}


PARTNERSHIP_SEED = ROOT / 'partnership_seed.json'
PARTNERSHIP_TEMPLATE = ROOT / 'partnership_report_template.html'

PARTNERSHIP_FIELDS = {
    'source_key','partenaire','raison_sociale','type_partenariat','statut_convention','signature_statut',
    'date_signature','date_debut','date_fin','duree_texte','renouvellement_texte','preavis_texte','objet',
    'expertise','exclusif','perimetre_exclusivite','dispositifs','programmes','representant','referent_interne',
    'secteur','source_referent','montants','facturation_texte','paiement_texte','resume_operationnel',
    'droits_partenaire','obligations_partenaire','points_vigilance','source_fichiers','source_reference',
    'clauses_source','niveau_confiance','notes_internes'
}
ENGAGEMENT_FIELDS = {
    'partie','service_code','categorie','titre','description','quantite_cible','unite','frequence','echeance',
    'obligatoire','priorite','statut','progression','responsable','preuve','commentaire_suivi','source_reference',
    'classification','dispositifs'
}
ACTION_FIELDS = {
    'engagement_id','service_code','type_action','titre','description','echeance','frequence','priorite','statut','progression',
    'responsable','preuve_attendue','preuve','commentaire_suivi','origine','ordre',
    'date_planification','heure_planification','duree_minutes','lieu','notes_calendrier'
}
REALISATION_FIELDS = {'date_prevue','date_realisation','quantite','statut','commentaire','preuve'}


def _filter_fields(data, allowed):
    return {k:v for k,v in (data or {}).items() if k in allowed}


def _partnership_slug(value):
    v=unicodedata.normalize('NFKD',str(value or '')).encode('ascii','ignore').decode('ascii').lower()
    v=re.sub(r'[^a-z0-9]+','-',v).strip('-') or 'partenaire'
    return v[:48]


def _default_action_from_engagement(e, partenariat_id):
    """Crée une action de pilotage générique pour un nouvel engagement saisi manuellement.
    Cette action est explicitement interne et ne modifie pas le contenu contractuel."""
    party=str(e.get('partie') or 'eurasante')
    title=str(e.get('titre') or 'Engagement')
    service=str(e.get('service_code') or '')
    if party=='partenaire':
        action_title=f"Suivre et relancer le partenaire — {title}"
        desc='Vérifier la réalisation de cet engagement par le partenaire, relancer si nécessaire et archiver une preuve exploitable.'
    elif party=='mutuel':
        action_title=f"Coordonner et tracer — {title}"
        desc='Coordonner les interlocuteurs des deux parties, planifier les étapes et documenter la réalisation ainsi que les éventuels écarts.'
    else:
        action_title=f"Planifier, réaliser et tracer — {title}"
        desc='Désigner le responsable, définir le prochain jalon, réaliser l’action attendue et conserver une preuve permettant de démontrer le respect de l’engagement.'
    proofs={
        'COMMUNICATION':'URL, capture d’écran, exemplaire du support ou validation de publication',
        'EVENEMENTIEL':'invitation / calendrier, liste de participants, compte rendu ou photo',
        'INCUBATION':'e-mail de mise en relation, invitation, compte rendu de rendez-vous ou comité',
        'FINANCE':'facture, bon de commande, preuve d’envoi et/ou preuve d’encaissement',
        'JURIDIQUE':'version signée, avenant, décision interne, courrier ou preuve de notification',
        'PARTENARIATS':'compte rendu de revue, e-mail de validation, tableau de suivi ou décision',
        'EURALIMENTAIRE':'invitation, compte rendu, feuille de présence, preuve de mentoring/jury',
        'ILAB':'convocation, support, feuille de présence ou compte rendu',
        'PARC':'offre publiée, communication aux résidents, URL ou justificatif de référencement'
    }
    return {
        'partenariat_id':partenariat_id,'engagement_id':e.get('id'),
        'source_key':f"{e.get('source_key')}-a01",'service_code':e.get('service_code'),
        'type_action':'suivi','titre':action_title,'description':desc,
        'echeance':e.get('echeance'),'frequence':e.get('frequence') or 'selon l’engagement',
        'priorite':e.get('priorite') or 'normale','statut':'a_faire','progression':0,
        'responsable':e.get('responsable') or '','preuve_attendue':proofs.get(service,'e-mail, compte rendu, justificatif ou autre preuve vérifiable'),
        'preuve':'','commentaire_suivi':'','origine':'pilotage_derive','ordre':10
    }



PARTNERSHIP_IMPORT_SCHEMA = 'eurasante.partnership.v5'
ENGAGEMENT_STATUSES = {'a_faire','en_cours','realise','non_applicable','alerte'}
ACTION_STATUSES = {'a_faire','en_cours','realise','bloque','non_applicable'}
PARTIES = {'eurasante','partenaire','mutuel'}


def _v5_action_columns_ready():
    """Teste la presence des colonnes de planification V5 sans modifier la base."""
    try:
        STATE.client.table('actions_pilotage_partenariat').select(
            'date_planification,heure_planification,duree_minutes,lieu,notes_calendrier'
        ).limit(1).execute()
        return True
    except Exception as exc:
        msg=str(exc).lower()
        if 'does not exist' in msg or 'schema cache' in msg or 'pgrst' in msg or 'column' in msg:
            return False
        raise


def _iso_date_or_none(value, field, errors):
    if value in (None,''):
        return None
    v=str(value).strip()
    try:
        return datetime.strptime(v,'%Y-%m-%d').date().isoformat()
    except Exception:
        errors.append(f'{field}: date attendue au format AAAA-MM-JJ ({v})')
        return None


def _time_or_none(value, field, errors):
    if value in (None,''):
        return None
    v=str(value).strip()[:5]
    try:
        datetime.strptime(v,'%H:%M')
        return v
    except Exception:
        errors.append(f'{field}: heure attendue au format HH:MM ({value})')
        return None


def _source_key(value, fallback=''):
    raw=str(value or fallback or '').strip()
    raw=unicodedata.normalize('NFKD',raw).encode('ascii','ignore').decode('ascii').lower()
    raw=re.sub(r'[^a-z0-9._-]+','-',raw).strip('-')
    return raw[:160]


def _bounded_int(value, default, minimum, maximum, field, errors):
    if value in (None,''):
        return default
    try:
        return max(minimum,min(maximum,int(value)))
    except Exception:
        errors.append(f'{field} doit etre un entier entre {minimum} et {maximum}.')
        return default


def _normalise_import_package(raw):
    """Valide et normalise le JSON d'une convention. Aucune ecriture ici."""
    errors=[]; warnings=[]
    if not isinstance(raw,dict):
        return None, ['Le JSON racine doit etre un objet.'], []
    schema=str(raw.get('schema_version') or raw.get('schema') or '').strip()
    if isinstance(raw.get('meta'),dict) and raw.get('meta',{}).get('template') is True:
        errors.append('Ce fichier est le modele V5: remplacer les valeurs d exemple et retirer meta.template avant import.')
    if schema and schema not in {PARTNERSHIP_IMPORT_SCHEMA,'5','v5','V5'}:
        warnings.append(f'Version de schema non standard: {schema}. Le format V5 sera applique.')
    part=raw.get('partenariat') or raw.get('partnership')
    if not isinstance(part,dict):
        return None, ['Objet "partenariat" manquant.'], warnings
    name=str(part.get('partenaire') or part.get('raison_sociale') or '').strip()
    if not name:
        errors.append('partenariat.partenaire est obligatoire.')
    pkey=_source_key(part.get('source_key'), f'{name}-{date.today().year}')
    if not pkey:
        errors.append('partenariat.source_key est obligatoire ou doit pouvoir etre derive du nom.')
    part=dict(part); part['partenaire']=name; part['source_key']=pkey
    for f in ('date_signature','date_debut','date_fin'):
        if f in part: part[f]=_iso_date_or_none(part.get(f),f'partenariat.{f}',errors)
    for f in ('dispositifs','programmes','droits_partenaire','obligations_partenaire','points_vigilance','source_fichiers'):
        if f in part and not isinstance(part[f],list):
            part[f]=[x.strip() for x in str(part[f]).split(',') if x.strip()]
    if 'montants' in part and not isinstance(part.get('montants'),list):
        errors.append('partenariat.montants doit etre un tableau JSON.')
    if 'clauses_source' in part and not isinstance(part.get('clauses_source'),dict):
        errors.append('partenariat.clauses_source doit etre un objet JSON.')

    services=[]; service_codes=set()
    for i,row in enumerate(raw.get('services') or []):
        if not isinstance(row,dict):
            errors.append(f'services[{i}] doit etre un objet.'); continue
        code=str(row.get('code') or '').strip().upper()
        nom=str(row.get('nom') or '').strip()
        if not code or not nom:
            errors.append(f'services[{i}] requiert code et nom.'); continue
        services.append({'code':code,'nom':nom,'description':str(row.get('description') or ''),'ordre':_bounded_int(row.get('ordre'),100,0,100000,f'services[{i}].ordre',errors),'actif':row.get('actif',True) is not False})
        service_codes.add(code)

    engagements=[]; ekeys=set()
    for i,row in enumerate(raw.get('engagements') or []):
        if not isinstance(row,dict):
            errors.append(f'engagements[{i}] doit etre un objet.'); continue
        e=dict(row)
        title=str(e.get('titre') or '').strip()
        key=_source_key(e.get('source_key'),f'{pkey}-eng-{i+1:03d}')
        if not title: errors.append(f'engagements[{i}].titre est obligatoire.')
        if key in ekeys: errors.append(f'engagements[{i}].source_key duplique: {key}')
        ekeys.add(key); e['source_key']=key; e['partenaire_source_key']=pkey
        party=str(e.get('partie') or 'eurasante').strip().lower()
        if party not in PARTIES:
            errors.append(f'engagements[{i}].partie invalide: {party}'); party='eurasante'
        e['partie']=party
        st=str(e.get('statut') or 'a_faire').strip()
        if st not in ENGAGEMENT_STATUSES:
            warnings.append(f'engagements[{i}].statut "{st}" remplace par a_faire.'); st='a_faire'
        e['statut']=st
        e['progression']=_bounded_int(e.get('progression'),100 if st=='realise' else 0,0,100,f'engagements[{i}].progression',errors)
        e['echeance']=_iso_date_or_none(e.get('echeance'),f'engagements[{i}].echeance',errors)
        e.setdefault('classification','clause'); e.setdefault('obligatoire',True); e.setdefault('priorite','normale')
        if e.get('service_code'): e['service_code']=str(e['service_code']).strip().upper()
        if 'dispositifs' in e and not isinstance(e['dispositifs'],list): e['dispositifs']=[x.strip() for x in str(e['dispositifs']).split(',') if x.strip()]
        engagements.append(e)

    actions=[]; akeys=set()
    for i,row in enumerate(raw.get('actions') or raw.get('taches') or []):
        if not isinstance(row,dict):
            errors.append(f'actions[{i}] doit etre un objet.'); continue
        a=dict(row)
        title=str(a.get('titre') or '').strip()
        key=_source_key(a.get('source_key'),f'{pkey}-act-{i+1:03d}')
        if not title: errors.append(f'actions[{i}].titre est obligatoire.')
        if key in akeys: errors.append(f'actions[{i}].source_key duplique: {key}')
        akeys.add(key); a['source_key']=key; a['partenaire_source_key']=pkey
        st=str(a.get('statut') or 'a_faire').strip()
        if st not in ACTION_STATUSES:
            warnings.append(f'actions[{i}].statut "{st}" remplace par a_faire.'); st='a_faire'
        a['statut']=st
        a['progression']=_bounded_int(a.get('progression'),100 if st=='realise' else 0,0,100,f'actions[{i}].progression',errors)
        a['echeance']=_iso_date_or_none(a.get('echeance'),f'actions[{i}].echeance',errors)
        a['date_planification']=_iso_date_or_none(a.get('date_planification'),f'actions[{i}].date_planification',errors)
        a['heure_planification']=_time_or_none(a.get('heure_planification'),f'actions[{i}].heure_planification',errors)
        if a.get('duree_minutes') not in (None,''):
            a['duree_minutes']=_bounded_int(a.get('duree_minutes'),60,5,1440,f'actions[{i}].duree_minutes',errors)
        else: a['duree_minutes']=60
        if a.get('service_code'): a['service_code']=str(a['service_code']).strip().upper()
        a['engagement_source_key']=_source_key(a.get('engagement_source_key'))
        a.setdefault('type_action','suivi'); a.setdefault('priorite','normale'); a.setdefault('origine','import_json_convention'); a.setdefault('ordre',10)
        actions.append(a)

    return {'schema_version':PARTNERSHIP_IMPORT_SCHEMA,'meta':raw.get('meta') or {},'services':services,'partenariat':part,'engagements':engagements,'actions':actions}, errors, warnings


def preview_partnership_import(raw):
    pack,errors,warnings=_normalise_import_package(raw)
    if not pack:
        return {'ok':False,'errors':errors,'warnings':warnings}
    current_services,_=safe_fetch_all('referentiel_services_partenariats')
    current_parts,_=safe_fetch_all('partenariats')
    current_eng,_=safe_fetch_all('engagements_partenariat')
    current_actions,actions_ready=safe_fetch_all('actions_pilotage_partenariat')
    v5_ready=_v5_action_columns_ready() if actions_ready else False
    svc={str(x.get('code')) for x in current_services}|{x['code'] for x in pack['services']}
    for i,e in enumerate(pack['engagements']):
        if e.get('service_code') and e['service_code'] not in svc: errors.append(f'engagements[{i}].service_code inconnu: {e["service_code"]}')
    for i,a in enumerate(pack['actions']):
        if a.get('service_code') and a['service_code'] not in svc: errors.append(f'actions[{i}].service_code inconnu: {a["service_code"]}')
        if a.get('engagement_source_key') and a['engagement_source_key'] not in {e['source_key'] for e in pack['engagements']}|{str(e.get('source_key')) for e in current_eng}:
            errors.append(f'actions[{i}].engagement_source_key introuvable: {a["engagement_source_key"]}')
    existing_p={str(x.get('source_key')) for x in current_parts}
    existing_e={str(x.get('source_key')) for x in current_eng}
    existing_a={str(x.get('source_key')) for x in current_actions}
    existing_s={str(x.get('code')) for x in current_services}
    if not actions_ready: errors.append('Table actions_pilotage_partenariat absente: executer la migration V4.1.')
    if actions_ready and not v5_ready and pack['actions']:
        errors.append('Import JSON V5 avec taches: executer MIGRATION_PARTENARIATS_V5.sql avant import. La migration est additive.')
    return {
        'ok':not errors,'errors':errors,'warnings':warnings,'v5_ready':v5_ready,
        'summary':{
            'partnership_create':0 if pack['partenariat']['source_key'] in existing_p else 1,
            'partnership_skip':1 if pack['partenariat']['source_key'] in existing_p else 0,
            'services_create':sum(1 for x in pack['services'] if x['code'] not in existing_s),
            'services_skip':sum(1 for x in pack['services'] if x['code'] in existing_s),
            'engagements_create':sum(1 for x in pack['engagements'] if x['source_key'] not in existing_e),
            'engagements_skip':sum(1 for x in pack['engagements'] if x['source_key'] in existing_e),
            'actions_create':sum(1 for x in pack['actions'] if x['source_key'] not in existing_a),
            'actions_skip':sum(1 for x in pack['actions'] if x['source_key'] in existing_a),
        },
        'normalized':pack
    }


def import_partnership_package(raw):
    preview=preview_partnership_import(raw)
    if not preview.get('ok'):
        raise ValueError('Import refuse: '+' | '.join(preview.get('errors') or ['format invalide']))
    pack=preview['normalized']; counts={'services_added':0,'partnerships_added':0,'partnerships_skipped':0,'engagements_added':0,'engagements_skipped':0,'actions_added':0,'actions_skipped':0}
    current_services,_=safe_fetch_all('referentiel_services_partenariats')
    existing_s={str(x.get('code')) for x in current_services}
    for row in pack['services']:
        if row['code'] in existing_s: continue
        STATE.client.table('referentiel_services_partenariats').insert(row).execute(); existing_s.add(row['code']); counts['services_added']+=1
    current_parts,_=safe_fetch_all('partenariats')
    existing_p={str(x.get('source_key')):x for x in current_parts}
    pkey=pack['partenariat']['source_key']
    if pkey in existing_p:
        partnership=existing_p[pkey]; counts['partnerships_skipped']=1
    else:
        payload=_filter_fields(pack['partenariat'],PARTNERSHIP_FIELDS)
        partnership=STATE.client.table('partenariats').insert(payload).execute().data[0]
        existing_p[pkey]=partnership; counts['partnerships_added']=1
    current_eng,_=safe_fetch_all('engagements_partenariat')
    engagement_by_source={str(x.get('source_key')):x for x in current_eng}
    for row in pack['engagements']:
        key=row['source_key']
        if key in engagement_by_source: counts['engagements_skipped']+=1; continue
        payload=_filter_fields(row,ENGAGEMENT_FIELDS); payload['source_key']=key; payload['partenariat_id']=partnership['id']
        created=STATE.client.table('engagements_partenariat').insert(payload).execute().data[0]
        engagement_by_source[key]=created; counts['engagements_added']+=1
    current_actions,_=safe_fetch_all('actions_pilotage_partenariat')
    existing_a={str(x.get('source_key')) for x in current_actions}
    for row in pack['actions']:
        key=row['source_key']
        if key in existing_a: counts['actions_skipped']+=1; continue
        payload=_filter_fields(row,ACTION_FIELDS); payload['source_key']=key; payload['partenariat_id']=partnership['id']
        ekey=row.get('engagement_source_key')
        if ekey and ekey in engagement_by_source: payload['engagement_id']=engagement_by_source[ekey]['id']
        STATE.client.table('actions_pilotage_partenariat').insert(payload).execute(); existing_a.add(key); counts['actions_added']+=1
    STATE.history(type_objet='partenariat',objet_id=partnership.get('id'),objet_nom=partnership.get('partenaire'),action='import JSON V5 non destructif',nouvelle_valeur=json.dumps(counts,ensure_ascii=False))
    return {'ok':True,**counts,'warnings':preview.get('warnings') or [],'partnership_id':partnership.get('id'),'partnership_source_key':pkey}


def _ics_escape(value):
    return str(value or '').replace('\\','\\\\').replace('\r','').replace('\n','\\n').replace(';','\\;').replace(',','\\,')


def build_partnership_actions_ics(action_ids):
    ids={str(x) for x in (action_ids or []) if x}
    if not ids: raise ValueError('Selectionner au moins une action.')
    actions,_=safe_fetch_all('actions_pilotage_partenariat')
    selected=[a for a in actions if str(a.get('id')) in ids]
    if not selected: raise ValueError('Aucune action selectionnee n a ete trouvee.')
    parts,_=safe_fetch_all('partenariats'); engs,_=safe_fetch_all('engagements_partenariat'); services,_=safe_fetch_all('referentiel_services_partenariats')
    pmap={str(x.get('id')):x for x in parts}; emap={str(x.get('id')):x for x in engs}; smap={str(x.get('code')):x for x in services}
    missing=[a.get('titre') for a in selected if not (a.get('date_planification') or a.get('echeance'))]
    if missing:
        shown=', '.join(str(x) for x in missing[:5]); more=' …' if len(missing)>5 else ''
        raise ValueError(f'{len(missing)} action(s) sans date de planification ni echeance: {shown}{more}')
    now=datetime.now(tz=ZoneInfo('UTC')).strftime('%Y%m%dT%H%M%SZ')
    lines=['BEGIN:VCALENDAR','VERSION:2.0','PRODID:-//Eurasante//CRM Partenariats V5//FR','CALSCALE:GREGORIAN','METHOD:PUBLISH','X-WR-CALNAME:Eurasante - Taches partenariats']
    for a in selected:
        p=pmap.get(str(a.get('partenariat_id')),{}) ; e=emap.get(str(a.get('engagement_id')),{}) ; svc=smap.get(str(a.get('service_code')),{})
        day=str(a.get('date_planification') or a.get('echeance'))[:10]
        hour=str(a.get('heure_planification') or '')[:5]
        title=f"[{p.get('partenaire') or 'Partenariat'}] {a.get('titre') or 'Action'}"
        desc=[f"Service: {svc.get('nom') or a.get('service_code') or 'Non affecte'}"]
        if a.get('responsable'): desc.append(f"Responsable: {a.get('responsable')}")
        if e.get('titre'): desc.append(f"Engagement: {e.get('titre')}")
        if a.get('description'): desc.append(f"Action: {a.get('description')}")
        if a.get('preuve_attendue'): desc.append(f"Preuve attendue: {a.get('preuve_attendue')}")
        if a.get('notes_calendrier'): desc.append(f"Note calendrier: {a.get('notes_calendrier')}")
        if a.get('commentaire_suivi'): desc.append(f"Suivi: {a.get('commentaire_suivi')}")
        if e.get('source_reference'): desc.append(f"Source convention: {e.get('source_reference')}")
        uid=f"{_source_key(a.get('source_key'),a.get('id'))}@eurasante-crm"
        lines += ['BEGIN:VEVENT',f'UID:{_ics_escape(uid)}',f'DTSTAMP:{now}',f'SUMMARY:{_ics_escape(title)}']
        if hour:
            try:
                dt=datetime.strptime(day+' '+hour,'%Y-%m-%d %H:%M')
                dur=max(5,int(a.get('duree_minutes') or 60))
                end=dt+timedelta(minutes=dur)
                lines += [f'DTSTART;TZID=Europe/Paris:{dt.strftime("%Y%m%dT%H%M%S")}',f'DTEND;TZID=Europe/Paris:{end.strftime("%Y%m%dT%H%M%S")}']
            except Exception:
                d=datetime.strptime(day,'%Y-%m-%d').date(); lines += [f'DTSTART;VALUE=DATE:{d.strftime("%Y%m%d")}',f'DTEND;VALUE=DATE:{(d+timedelta(days=1)).strftime("%Y%m%d")}']
        else:
            d=datetime.strptime(day,'%Y-%m-%d').date(); lines += [f'DTSTART;VALUE=DATE:{d.strftime("%Y%m%d")}',f'DTEND;VALUE=DATE:{(d+timedelta(days=1)).strftime("%Y%m%d")}']
        if a.get('lieu'): lines.append(f'LOCATION:{_ics_escape(a.get("lieu"))}')
        lines += [f'DESCRIPTION:{_ics_escape(chr(10).join(desc))}',f'CATEGORIES:{_ics_escape("EURASANTE,PARTENARIAT,"+str(a.get("service_code") or "AUTRE"))}','STATUS:CONFIRMED','X-MICROSOFT-CDO-BUSYSTATUS:BUSY','END:VEVENT']
    lines.append('END:VCALENDAR')
    return ('\r\n'.join(lines)+'\r\n').encode('utf-8')

def partnership_payload():
    partnerships, ready = safe_fetch_all('partenariats','partenaire')
    if not ready:
        return {'ready':False,'actions_ready':False,'v5_ready':False,'services':[],'partenariats':[],'engagements':[],'actions':[],'realisations':[],
                'migration':'MIGRATION_PARTENARIATS_V4.sql'}
    services,_=safe_fetch_all('referentiel_services_partenariats','ordre')
    engagements,_=safe_fetch_all('engagements_partenariat','created_at')
    actions, actions_ready=safe_fetch_all('actions_pilotage_partenariat','created_at')
    realisations,_=safe_fetch_all('realisations_engagement_partenariat','created_at')
    source_by_id={str(x.get('id')):x.get('source_key') for x in partnerships}
    engagement_source_by_id={str(x.get('id')):x.get('source_key') for x in engagements}
    for e in engagements:
        e['partenaire_source_key']=source_by_id.get(str(e.get('partenariat_id')),'')
    for a in actions:
        a['partenaire_source_key']=source_by_id.get(str(a.get('partenariat_id')),'')
        a['engagement_source_key']=engagement_source_by_id.get(str(a.get('engagement_id')),'')
    return {'ready':True,'actions_ready':actions_ready,'v5_ready':(_v5_action_columns_ready() if actions_ready else False),'services':services,'partenariats':partnerships,'engagements':engagements,
            'actions':actions,'realisations':realisations,'meta':{'generated_date':date.today().isoformat(),
            'warning':'Les engagements décrivent les obligations ; les actions de pilotage traduisent les contrôles et tâches internes nécessaires pour en sécuriser le respect. La convention source demeure prioritaire.'}}


def seed_partnerships():
    if not PARTNERSHIP_SEED.exists():
        raise ValueError('Référentiel partnership_seed.json introuvable.')
    # Vérifie d'abord que la migration a bien été appliquée.
    current,ready=safe_fetch_all('partenariats')
    if not ready:
        raise ValueError('Module Partenariats non initialisé. Exécuter MIGRATION_PARTENARIATS_V4.sql dans Supabase.')
    seed=json.loads(PARTNERSHIP_SEED.read_text(encoding='utf-8'))
    service_added=part_added=eng_added=0
    existing_services={str(x.get('code')) for x in safe_fetch_all('referentiel_services_partenariats')[0]}
    for row in seed.get('services',[]):
        if row.get('code') not in existing_services:
            STATE.client.table('referentiel_services_partenariats').insert({k:v for k,v in row.items() if k in {'code','nom','description','ordre','actif'}}).execute()
            existing_services.add(row.get('code')); service_added+=1
    existing={str(x.get('source_key')):x for x in current}
    for row in seed.get('partenariats',[]):
        key=str(row.get('source_key') or '')
        if not key: continue
        if key not in existing:
            payload=_filter_fields(row,PARTNERSHIP_FIELDS)
            created=STATE.client.table('partenariats').insert(payload).execute().data[0]
            existing[key]=created; part_added+=1
    existing_eng={str(x.get('source_key')) for x in safe_fetch_all('engagements_partenariat')[0]}
    for row in seed.get('engagements',[]):
        key=str(row.get('source_key') or ''); pkey=str(row.get('partenaire_source_key') or '')
        if not key or key in existing_eng or pkey not in existing: continue
        payload=_filter_fields(row,ENGAGEMENT_FIELDS)
        payload['source_key']=key; payload['partenariat_id']=existing[pkey]['id']
        STATE.client.table('engagements_partenariat').insert(payload).execute()
        existing_eng.add(key); eng_added+=1
    # Recharge les engagements pour disposer des identifiants, y compris pour les lignes déjà existantes.
    engagement_rows,_=safe_fetch_all('engagements_partenariat')
    engagement_by_source={str(x.get('source_key')):x for x in engagement_rows}
    action_added=0
    action_rows,actions_ready=safe_fetch_all('actions_pilotage_partenariat')
    if actions_ready:
        existing_actions={str(x.get('source_key')) for x in action_rows}
        for row in seed.get('actions',[]):
            key=str(row.get('source_key') or ''); pkey=str(row.get('partenaire_source_key') or '')
            ekey=str(row.get('engagement_source_key') or '')
            if not key or key in existing_actions or pkey not in existing: continue
            payload=_filter_fields(row,ACTION_FIELDS)
            payload['source_key']=key; payload['partenariat_id']=existing[pkey]['id']
            if ekey and ekey in engagement_by_source:
                payload['engagement_id']=engagement_by_source[ekey]['id']
            STATE.client.table('actions_pilotage_partenariat').insert(payload).execute()
            existing_actions.add(key); action_added+=1
    STATE.history(type_objet='partenariat',objet_nom='Référentiel conventions',action='initialisation non destructive',nouvelle_valeur=f'{part_added} partenariats, {eng_added} engagements, {action_added} actions, {service_added} services ajoutés')
    return {'services_added':service_added,'partnerships_added':part_added,'engagements_added':eng_added,'actions_added':action_added,
            'actions_ready':actions_ready,
            'message':'Import non destructif terminé. Les lignes déjà présentes n’ont pas été modifiées.' + ('' if actions_ready else ' Exécuter MIGRATION_PARTENARIATS_V41_ACTIONS.sql pour activer les actions de pilotage.')}


def build_partnership_html(data=None):
    if not PARTNERSHIP_TEMPLATE.exists():
        raise ValueError('Template HTML partenariats introuvable.')
    data=data or partnership_payload()
    if not data.get('ready',True):
        raise ValueError('Module Partenariats non initialisé dans Supabase.')
    if data.get('actions_ready') is False:
        raise ValueError('Actions de pilotage non initialisées. Exécuter MIGRATION_PARTENARIATS_V41_ACTIONS.sql puis recharger le référentiel avant de générer le HTML.')
    # Le fichier autonome n'a pas besoin des identifiants techniques DB.
    clean={'meta':data.get('meta') or {},'services':data.get('services') or [],
           'partenariats':data.get('partenariats') or [],'engagements':data.get('engagements') or [],
           'actions':data.get('actions') or []}
    raw=json.dumps(clean,ensure_ascii=False,default=str).replace('</','<\\/')
    return PARTNERSHIP_TEMPLATE.read_text(encoding='utf-8').replace('__DATA_JSON__',raw).encode('utf-8')

def _relation_domain(value):
    value=str(value or '').strip().lower()
    if '@' in value:
        value=value.rsplit('@',1)[-1]
    return value.rstrip('.')

def _relation_excluded(value):
    d=_relation_domain(value)
    return d in RELATION_EXCLUDED_DOMAINS or any(d.endswith('.'+x) for x in RELATION_EXCLUDED_DOMAINS)

def relation_payload():
    organisations, ready = safe_fetch_all('organisations','nom')
    if not ready:
        return {'ready':False,'organisations':[],'contacts':[],'rdv':[],'rdv_contacts':[]}
    # Masquage défensif : même si d'anciennes lignes existent encore en base,
    # elles ne polluent plus le référentiel affiché.
    excluded_org_ids={str(x.get('id')) for x in organisations if _relation_excluded(x.get('domaine')) or _relation_excluded(x.get('nom'))}
    organisations=[x for x in organisations if str(x.get('id')) not in excluded_org_ids]
    contacts,_=safe_fetch_all('contacts_relation','email')
    contacts=[x for x in contacts if str(x.get('organisation_id') or '') not in excluded_org_ids and not _relation_excluded(x.get('email'))]
    allowed_contact_ids={str(x.get('id')) for x in contacts}
    rdv,_=safe_fetch_all('rendez_vous','date_rdv')
    rdv=[x for x in rdv if str(x.get('organisation_id') or '') not in excluded_org_ids]
    allowed_rdv_ids={str(x.get('id')) for x in rdv}
    links,_=safe_fetch_all('rdv_contacts')
    links=[x for x in links if str(x.get('rdv_id')) in allowed_rdv_ids and str(x.get('contact_id')) in allowed_contact_ids]
    return {'ready':True,'organisations':organisations,'contacts':contacts,'rdv':rdv,'rdv_contacts':links}

def _relation_norm_text(value):
    value=unicodedata.normalize('NFKD', str(value or ''))
    value=''.join(c for c in value if not unicodedata.combining(c)).lower()
    return re.sub(r'\s+',' ',value).strip()

# Exclusions volontairement conservatrices. Elles ne s'appliquent que lorsqu'aucun
# contact/domaine externe exploitable n'est présent. Un RDV avec Pfizer + Clubster,
# par exemple, est toujours conservé et rattaché à Pfizer.
RELATION_STRONG_IGNORE_PATTERNS=[
    r'^tt$', r'^rtt$', r'^absence\b', r'^perso(?:nnel)?$', r'^vacances scolaires',
    r'^permis$', r'^retrait$', r'^chaudiere$', r'^controle technique',
    r'^recup\b', r'^sur site$', r'^inauguration$', r'^teambuilding\b',
    r'^soiree de fin d.?annee$', r'^pot de depart\b', r'^debrief codir\b',
    r'^webinaire interne\b', r'^presentation process crm commercial partenariat$',
    r'^presentation strategie ia attractivite partenariat$'
]
RELATION_COMMERCIAL_HINTS=(
    'rdv','rendez-vous','rendez vous','call','visio','echange','échange','rencontre',
    'partenariat','prospect','offre','relance','client','sponsor','fondation','banque',
    'assurance','mutuelle','eurasante x','eurasante *','eurasanté x','eurasanté *'
)
GENERIC_MAIL_DOMAINS={
    'gmail.com','googlemail.com','hotmail.com','outlook.com','live.fr','live.com',
    'icloud.com','me.com','yahoo.com','yahoo.fr','orange.fr','wanadoo.fr'
}

def _relation_should_ignore_without_external(event):
    title=_relation_norm_text(event.get('title'))
    if not title:
        return True,'objet vide'
    for pat in RELATION_STRONG_IGNORE_PATTERNS:
        if re.search(pat,title,re.I):
            return True,'interne/personnel explicite'
    # Les événements sans domaine/contact mais dont l'objet comporte un signal
    # commercial sont conservés en file « À qualifier ».
    if any(h in title for h in RELATION_COMMERCIAL_HINTS):
        return False,''
    # Un objet court mais nominatif peut être un vrai échange (ex. "CIC", "Franck").
    # On le conserve sauf s'il correspond à une exclusion forte ci-dessus.
    if len(title) >= 3 and title not in {'option','test','rappel'}:
        return False,''
    return True,'aucun signal relationnel exploitable'

def import_relation_history():
    source=ROOT/'rdv_history.json'
    if not source.exists(): raise ValueError('Fichier rdv_history.json introuvable.')
    events=json.loads(source.read_text(encoding='utf-8')).get('events',[])
    inserted=updated=0
    org_created=contacts_created=0
    today=date.today().isoformat()
    stats={
        'analysed':len(events),'inserted':0,'updated':0,'duplicates_updated':0,
        'with_organisation':0,'without_organisation':0,'ambiguous_organisation':0,
        'contacts_created':0,'organisations_created':0,'ignored':0,'errors':0,
        'cancelled':0,'planned':0,'to_qualify':0
    }
    ignored_reasons={}
    ignored_examples=[]
    error_examples=[]

    for e in events:
        uid=str(e.get('uid') or '').strip()
        if not uid:
            stats['ignored']+=1
            ignored_reasons['UID calendrier absent']=ignored_reasons.get('UID calendrier absent',0)+1
            continue
        try:
            raw_domains=[]
            for x in (e.get('domains') or []):
                d=_relation_domain(x)
                if d and 'resource.calendar.google.com' not in d and d not in raw_domains:
                    raw_domains.append(d)
            domains=[d for d in raw_domains if not _relation_excluded(d)]

            external_emails=[]
            for x in (e.get('external') or []):
                email=str(x or '').strip().lower()
                if not email or email.endswith('@resource.calendar.google.com') or _relation_excluded(email):
                    continue
                if email not in external_emails:
                    external_emails.append(email)

            # Si l'analyse n'a trouvé aucun domaine mais qu'un email externe existe,
            # on récupère son domaine. Les domaines génériques restent des contacts,
            # mais ne créent pas automatiquement une organisation.
            email_domains=[]
            for email in external_emails:
                d=_relation_domain(email)
                if d and d not in email_domains and not _relation_excluded(d):
                    email_domains.append(d)
            candidate_domains=[]
            for d in domains+email_domains:
                if d not in candidate_domains and d not in GENERIC_MAIL_DOMAINS:
                    candidate_domains.append(d)

            has_external=bool(domains or external_emails)
            if not has_external:
                ignore,reason=_relation_should_ignore_without_external(e)
                if ignore:
                    stats['ignored']+=1
                    ignored_reasons[reason]=ignored_reasons.get(reason,0)+1
                    if len(ignored_examples)<25:
                        ignored_examples.append({'uid':uid,'date':e.get('date'),'objet':e.get('title'),'raison':reason})
                    continue

            domain=candidate_domains[0] if len(candidate_domains)==1 else None
            if len(candidate_domains)>1:
                stats['ambiguous_organisation']+=1
            org_id=None
            if domain:
                found=STATE.client.table('organisations').select('*').ilike('domaine',domain).execute().data or []
                if found:
                    org_id=found[0]['id']
                else:
                    row=STATE.client.table('organisations').insert({
                        'nom':domain,'domaine':domain,'statut_relation':'a_qualifier'
                    }).execute().data[0]
                    org_id=row['id']; org_created+=1; stats['organisations_created']+=1
                stats['with_organisation']+=1
            else:
                stats['without_organisation']+=1

            title=str(e.get('title') or '').strip() or 'Rendez-vous à qualifier'
            low=_relation_norm_text(title)
            if 'annul' in low or 'cancel' in low:
                statut='annule'; stats['cancelled']+=1
            elif str(e.get('date') or '')>today:
                statut='planifie'; stats['planned']+=1
            else:
                statut='a_qualifier'; stats['to_qualify']+=1

            payload={
                'uid_calendrier':uid,'date_rdv':e.get('date'),'objet':title,
                'lieu':e.get('location') or None,'organisateur':e.get('organizer') or None,
                'description':e.get('description') or None,'organisation_id':org_id,
                'organisation_source':e.get('org') or None,'statut':statut,
                'source':'calendrier_historique'
            }
            existing=STATE.client.table('rendez_vous').select('id').eq('uid_calendrier',uid).execute().data or []
            if existing:
                rid=existing[0]['id']
                STATE.client.table('rendez_vous').update(payload).eq('id',rid).execute()
                updated+=1; stats['updated']+=1; stats['duplicates_updated']+=1
            else:
                rid=STATE.client.table('rendez_vous').insert(payload).execute().data[0]['id']
                inserted+=1; stats['inserted']+=1

            for email in external_emails:
                c=STATE.client.table('contacts_relation').select('*').ilike('email',email).execute().data or []
                if c:
                    cid=c[0]['id']
                    # Si le contact existait sans organisation et que celle-ci est maintenant connue,
                    # on enrichit sans écraser une autre organisation déjà renseignée.
                    if org_id and not c[0].get('organisation_id'):
                        STATE.client.table('contacts_relation').update({'organisation_id':org_id}).eq('id',cid).execute()
                else:
                    cid=STATE.client.table('contacts_relation').insert({'organisation_id':org_id,'email':email}).execute().data[0]['id']
                    contacts_created+=1; stats['contacts_created']+=1
                link=STATE.client.table('rdv_contacts').select('*').eq('rdv_id',rid).eq('contact_id',cid).execute().data or []
                if not link:
                    STATE.client.table('rdv_contacts').insert({'rdv_id':rid,'contact_id':cid}).execute()
        except Exception as exc:
            stats['errors']+=1
            if len(error_examples)<15:
                error_examples.append({'uid':uid,'date':e.get('date'),'objet':e.get('title'),'erreur':str(exc)[:300]})
            # L'import continue : une ligne défectueuse ne bloque plus les 1 700 suivantes.
            continue

    STATE.history(
        type_objet='relations',objet_nom='Historique RDV',action='import non destructif exhaustif',
        nouvelle_valeur=f"{inserted} ajoutes, {updated} maj, {stats['ignored']} ignores, {stats['errors']} erreurs"
    )
    return {
        'inserted':inserted,'updated':updated,'organisations_created':org_created,
        'contacts_created':contacts_created,'stats':stats,
        'ignored_reasons':ignored_reasons,'ignored_examples':ignored_examples,
        'error_examples':error_examples
    }

class Handler(SimpleHTTPRequestHandler):
    def translate_path(self, path):
        rel = urlparse(path).path.lstrip('/') or 'app.html'
        return str(ROOT / rel)
    def log_message(self, fmt,*args): print('[HTTP]',fmt%args)
    def end_headers(self):
        self.send_header('Cache-Control','no-store, no-cache, must-revalidate')
        super().end_headers()
    def json_out(self,obj,status=200):
        b=json.dumps(obj,ensure_ascii=False,default=str).encode('utf-8')
        self.send_response(status); self.send_header('Content-Type','application/json; charset=utf-8'); self.send_header('Content-Length',str(len(b))); self.end_headers(); self.wfile.write(b)
    def body_json(self):
        n=int(self.headers.get('Content-Length','0')); return json.loads(self.rfile.read(n).decode('utf-8') or '{}')
    def fail(self,e):
        code=401 if isinstance(e,PermissionError) else 400
        self.json_out({"ok":False,"error":str(e)},code)

    def do_GET(self):
        try:
            p=urlparse(self.path); path=p.path
            if path=='/api/auth/status': return self.json_out({"authenticated":bool(STATE.email),"email":STATE.email})
            if path=='/api/config': return self.json_out({"version":CONFIG.get('version'),"refresh_seconds":REFRESH_SECONDS,"odj_names":ODJ_NAMES,"statuses":STATUSES})
            if path=='/api/data':
                STATE.require_auth()
                prospects=STATE.fetch_all('prospects','nom_sponsor')
                projects=STATE.fetch_all('projets','nom_projet')
                prospects=enrich_prospects_with_projects(prospects,projects)
                sourcing=STATE.fetch_all('sourcing','created_at')
                return self.json_out({"prospects":prospects,"projects":projects,"sourcing":sourcing,"relations":relation_payload(),"partnerships":partnership_payload(),"user":STATE.email,"server_time":datetime.now().isoformat()})
            if path=='/api/partnerships/import-template':
                STATE.require_auth()
                model_path=ROOT/'MODELE_IMPORT_CONVENTION_V5.json'
                if not model_path.exists(): raise ValueError('Modele JSON V5 introuvable.')
                payload=model_path.read_bytes()
                self.send_response(200); self.send_header('Content-Type','application/json; charset=utf-8'); self.send_header('Content-Disposition','attachment; filename=MODELE_IMPORT_CONVENTION_V5.json'); self.send_header('Content-Length',str(len(payload))); self.end_headers(); self.wfile.write(payload); return
            if path=='/api/partnerships/export-html':
                STATE.require_auth()
                payload=build_partnership_html()
                filename=f"Suivi_Partenariats_Eurasante_{date.today().isoformat()}.html"
                self.send_response(200)
                self.send_header('Content-Type','text/html; charset=utf-8')
                self.send_header('Content-Disposition',f'attachment; filename="{filename}"')
                self.send_header('Content-Length',str(len(payload)))
                self.end_headers(); self.wfile.write(payload); return
            if path=='/api/history':
                STATE.require_auth(); rows=STATE.fetch_all('historique','created_at'); return self.json_out(rows[-500:])
            if path=='/api/export/json':
                STATE.require_auth(); out={t:STATE.fetch_all(t) for t in ('prospects','projets','sourcing','historique')}; return self.json_out(out)
            if path=='/api/export/xlsx':
                STATE.require_auth()
                payload = build_xlsx_export()
                filename = f"Reporting_ODJ_export_{date.today().isoformat()}.xlsx"
                self.send_response(200)
                self.send_header('Content-Type','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                self.send_header('Content-Disposition',f'attachment; filename="{filename}"')
                self.send_header('Content-Length',str(len(payload)))
                self.end_headers()
                self.wfile.write(payload)
                return
            return super().do_GET()
        except Exception as e: return self.fail(e)

    def do_POST(self):
        try:
            p=urlparse(self.path); path=p.path
            if path=='/api/auth/login':
                d=self.body_json(); return self.json_out({"ok":True,**STATE.login(d.get('email','').strip(),d.get('password',''))})
            if path=='/api/auth/logout': STATE.logout(); return self.json_out({"ok":True})
            STATE.require_auth()
            if path=='/api/partnerships/import-json-preview':
                return self.json_out(preview_partnership_import(self.body_json()))
            if path=='/api/partnerships/import-json':
                return self.json_out(import_partnership_package(self.body_json()))
            if path=='/api/partnerships/export-ics':
                raw=self.body_json(); payload=build_partnership_actions_ics(raw.get('action_ids') or [])
                filename=f"Taches_Partenariats_Eurasante_{date.today().isoformat()}.ics"
                self.send_response(200); self.send_header('Content-Type','text/calendar; charset=utf-8'); self.send_header('Content-Disposition',f'attachment; filename="{filename}"'); self.send_header('Content-Length',str(len(payload))); self.end_headers(); self.wfile.write(payload); return
            if path=='/api/partnerships/seed':
                return self.json_out({'ok':True,**seed_partnerships()})
            if path=='/api/partnership/add':
                raw=self.body_json(); d=_filter_fields(raw,PARTNERSHIP_FIELDS)
                name=str(d.get('partenaire') or '').strip()
                if not name: raise ValueError('Nom du partenaire obligatoire')
                d['partenaire']=name
                if not d.get('source_key'):
                    d['source_key']=f"{_partnership_slug(name)}-{date.today().year}-{int(time.time())}"
                d.setdefault('statut_convention','active'); d.setdefault('signature_statut','à renseigner')
                r=STATE.client.table('partenariats').insert(d).execute().data[0]
                STATE.history(type_objet='partenariat',objet_id=r['id'],objet_nom=name,action='création')
                return self.json_out({'ok':True,'row':r})
            if path=='/api/partnership/update':
                raw=self.body_json(); rid=raw.pop('id',None)
                if not rid: raise ValueError('Identifiant partenariat manquant')
                d=_filter_fields(raw,PARTNERSHIP_FIELDS); d.pop('source_key',None)
                d['updated_at']=datetime.now().isoformat()
                r=STATE.client.table('partenariats').update(d).eq('id',rid).execute().data
                STATE.history(type_objet='partenariat',objet_id=rid,objet_nom=raw.get('partenaire',''),action='modification')
                return self.json_out({'ok':True,'row':(r or [{}])[0]})
            if path=='/api/partnership/engagement/add':
                raw=self.body_json(); pid=raw.pop('partenariat_id',None)
                if not pid: raise ValueError('Partenariat manquant')
                d=_filter_fields(raw,ENGAGEMENT_FIELDS)
                if not str(d.get('titre') or '').strip(): raise ValueError('Titre de l engagement obligatoire')
                d['partenariat_id']=pid
                d['source_key']=raw.get('source_key') or f"manual-{_partnership_slug(d.get('titre'))}-{int(time.time()*1000)}"
                r=STATE.client.table('engagements_partenariat').insert(d).execute().data[0]
                action=None
                try:
                    _,actions_ready=safe_fetch_all('actions_pilotage_partenariat')
                    if actions_ready:
                        ap=_default_action_from_engagement(r,pid)
                        action=STATE.client.table('actions_pilotage_partenariat').insert(ap).execute().data[0]
                except Exception:
                    # L'engagement reste créé même si la couche Actions n'est pas encore disponible.
                    action=None
                return self.json_out({'ok':True,'row':r,'action':action})
            if path=='/api/partnership/engagement/update':
                raw=self.body_json(); rid=raw.pop('id',None)
                if not rid: raise ValueError('Engagement manquant')
                d=_filter_fields(raw,ENGAGEMENT_FIELDS); d['updated_at']=datetime.now().isoformat()
                r=STATE.client.table('engagements_partenariat').update(d).eq('id',rid).execute().data
                return self.json_out({'ok':True,'row':(r or [{}])[0]})
            if path=='/api/partnership/action/add':
                raw=self.body_json(); pid=raw.pop('partenariat_id',None); eid=raw.pop('engagement_id',None)
                if not pid: raise ValueError('Partenariat manquant')
                d=_filter_fields(raw,ACTION_FIELDS)
                if not str(d.get('titre') or '').strip(): raise ValueError('Titre de l action obligatoire')
                d['partenariat_id']=pid
                if eid: d['engagement_id']=eid
                d['source_key']=raw.get('source_key') or f"manual-action-{_partnership_slug(d.get('titre'))}-{int(time.time()*1000)}"
                r=STATE.client.table('actions_pilotage_partenariat').insert(d).execute().data[0]
                return self.json_out({'ok':True,'row':r})
            if path=='/api/partnership/action/update':
                raw=self.body_json(); rid=raw.pop('id',None)
                if not rid: raise ValueError('Action manquante')
                d=_filter_fields(raw,ACTION_FIELDS); d['updated_at']=datetime.now().isoformat()
                r=STATE.client.table('actions_pilotage_partenariat').update(d).eq('id',rid).execute().data
                return self.json_out({'ok':True,'row':(r or [{}])[0]})
            if path=='/api/partnership/realisation/add':
                raw=self.body_json(); eid=raw.pop('engagement_id',None)
                if not eid: raise ValueError('Engagement manquant')
                d=_filter_fields(raw,REALISATION_FIELDS); d['engagement_id']=eid
                r=STATE.client.table('realisations_engagement_partenariat').insert(d).execute().data[0]
                return self.json_out({'ok':True,'row':r})
            if path=='/api/prospect/add':
                raw=self.body_json(); project_ids=_normalise_project_ids(raw.pop('projet_ids',None))
                d=clean_payload('prospects',raw)
                if project_ids is not None:
                    _validate_added_project_ids(project_ids,[])
                    d['projet_cible']=_primary_project_name(project_ids,d.get('projet_cible',''))
                else:
                    validate_new_project_attachment(d.get('projet_cible'))
                d.setdefault('date_maj',date.today().isoformat())
                r=STATE.client.table('prospects').insert(d).execute().data[0]
                if project_ids is not None:
                    sync_prospect_projects(r['id'],project_ids,d.get('projet_cible',''))
                    r=STATE.client.table('prospects').select('*').eq('id',r['id']).single().execute().data
                STATE.history(type_objet='prospect',objet_id=r['id'],objet_nom=r['nom_sponsor'],action='création')
                return self.json_out({"ok":True,"row":r})
            if path=='/api/prospect/update':
                raw=self.body_json(); rid=raw.pop('id'); project_ids=_normalise_project_ids(raw.pop('projet_ids',None))
                old=(STATE.client.table('prospects').select('*').eq('id',rid).single().execute().data or {})
                d=clean_payload('prospects',raw)
                if project_ids is None and 'projet_cible' in d and norm(d.get('projet_cible')) != norm(old.get('projet_cible')):
                    validate_new_project_attachment(d.get('projet_cible'))
                if project_ids is not None:
                    current_ids=[str(x.get('projet_id')) for x in _relations_for_prospect(rid) if x.get('projet_id')]
                    _validate_added_project_ids(project_ids,current_ids)
                    d['projet_cible']=_primary_project_name(project_ids,old.get('projet_cible',''))
                d.setdefault('date_maj',date.today().isoformat())
                r=STATE.client.table('prospects').update(d).eq('id',rid).execute().data[0]
                relation_change=None
                if project_ids is not None:
                    relation_change=sync_prospect_projects(rid,project_ids,old.get('projet_cible',''))
                    r=STATE.client.table('prospects').select('*').eq('id',rid).single().execute().data
                for k,v in d.items():
                    if str(old.get(k,''))!=str(v if v is not None else ''): STATE.history(type_objet='prospect',objet_id=rid,objet_nom=r.get('nom_sponsor'),action='modification',champ=k,ancienne_valeur=str(old.get(k,'') or ''),nouvelle_valeur=str(v or ''))
                if relation_change and (relation_change['added'] or relation_change['removed']):
                    STATE.history(type_objet='prospect',objet_id=rid,objet_nom=r.get('nom_sponsor'),action='modification',champ='projets associés',ancienne_valeur='',nouvelle_valeur=json.dumps(relation_change,ensure_ascii=False))
                return self.json_out({"ok":True,"row":r})
            if path=='/api/prospect/delete':
                d=self.body_json(); rid=d['id']; old=STATE.client.table('prospects').select('*').eq('id',rid).single().execute().data; STATE.client.table('prospects').delete().eq('id',rid).execute(); STATE.history(type_objet='prospect',objet_id=rid,objet_nom=(old or {}).get('nom_sponsor'),action='suppression'); return self.json_out({"ok":True})
            if path=='/api/project/add':
                d=clean_payload('projets',self.body_json()); d.setdefault('actif', True)
                r=STATE.client.table('projets').insert(d).execute().data[0]
                STATE.history(type_objet='projet',objet_id=r['id'],objet_nom=r['nom_projet'],action='création')
                return self.json_out({"ok":True,"row":r})
            if path=='/api/project/update':
                raw=self.body_json(); rid=raw.pop('id')
                old=STATE.client.table('projets').select('*').eq('id',rid).single().execute().data or {}
                d=clean_payload('projets',raw)
                rows=STATE.client.table('projets').update(d).eq('id',rid).execute().data or []
                # PostgREST peut renvoyer une représentation minimale selon la configuration : on relit toujours.
                r=STATE.client.table('projets').select('*').eq('id',rid).single().execute().data or {}
                if not r: raise RuntimeError('Projet introuvable après mise à jour')
                if norm(old.get('nom_projet')) != norm(r.get('nom_projet')) and old.get('nom_projet'):
                    # Maintient le champ historique projet_cible cohérent pour les exports XLSX.
                    STATE.client.table('prospects').update({'projet_cible':r.get('nom_projet')}).eq('projet_cible',old.get('nom_projet')).execute()
                if 'actif' in d and bool(r.get('actif')) != bool(d['actif']):
                    raise RuntimeError("La modification du statut Actif/Désactivé n'a pas été appliquée dans Supabase.")
                for k,v in d.items():
                    if str(old.get(k,''))!=str(v if v is not None else ''):
                        STATE.history(type_objet='projet',objet_id=rid,objet_nom=r.get('nom_projet'),action='modification',champ=k,ancienne_valeur=str(old.get(k,'') or ''),nouvelle_valeur=str(v or ''))
                return self.json_out({"ok":True,"row":r,"verified":True})
            if path=='/api/project/bulk-active':
                raw=self.body_json(); ids=raw.get('ids') or []
                requested=raw.get('actif')
                if not isinstance(requested, bool):
                    raise ValueError("Le statut actif doit être un booléen true/false")
                active=requested
                if not isinstance(ids,list) or not ids: raise ValueError('Aucun projet sélectionné')
                ids=list(dict.fromkeys(str(x) for x in ids if x))
                results=[]; failures=[]
                for rid in ids:
                    try:
                        old=STATE.client.table('projets').select('id,nom_projet,actif').eq('id',rid).single().execute().data or {}
                        if not old:
                            failures.append({'id':rid,'error':'Projet introuvable'}); continue
                        old_active=old.get('actif',True) is not False
                        STATE.client.table('projets').update({'actif':active}).eq('id',rid).execute()
                        check=STATE.client.table('projets').select('id,nom_projet,actif').eq('id',rid).single().execute().data or {}
                        if not check or bool(check.get('actif')) != active:
                            failures.append({'id':rid,'nom':old.get('nom_projet'),'error':'Valeur non confirmée après écriture'}); continue
                        results.append({'id':rid,'nom':check.get('nom_projet'),'actif':bool(check.get('actif'))})
                        if old_active!=active:
                            STATE.history(type_objet='projet',objet_id=rid,objet_nom=old.get('nom_projet'),action='modification en masse',champ='actif',ancienne_valeur='true' if old_active else 'false',nouvelle_valeur='true' if active else 'false')
                    except Exception as exc:
                        failures.append({'id':rid,'error':str(exc)})
                if failures:
                    return self.json_out({'ok':False,'updated':len(results),'requested':len(ids),'actif':active,'failures':failures},409)
                return self.json_out({'ok':True,'updated':len(results),'requested':len(ids),'actif':active,'rows':results,'verified':True})
            if path=='/api/sourcing/to-crm':
                d=self.body_json(); sid=d['id']; project_ids=_normalise_project_ids(d.get('projet_ids'))
                s=STATE.client.table('sourcing').select('*').eq('id',sid).single().execute().data
                if not s: raise ValueError('Cible sourcing introuvable')
                if project_ids is not None: _validate_added_project_ids(project_ids,[])
                else: validate_new_project_attachment(d.get('projet_cible'))
                existing=STATE.client.table('prospects').select('id,nom_sponsor,projet_cible').ilike('nom_sponsor',s['nom_cible']).execute().data or []
                if existing:
                    rid=existing[0]['id']
                    if project_ids:
                        current=[str(x.get('projet_id')) for x in _relations_for_prospect(rid) if x.get('projet_id')]
                        merged=list(dict.fromkeys(current+project_ids))
                        sync_prospect_projects(rid,merged,existing[0].get('projet_cible',''))
                else:
                    primary=_primary_project_name(project_ids or [],'') if project_ids is not None else (d.get('projet_cible') or None)
                    r=STATE.client.table('prospects').insert({'nom_sponsor':s['nom_cible'],'statut':'1. Prise de contact','date_maj':date.today().isoformat(),'projet_cible':primary}).execute().data[0]; rid=r['id']
                    if project_ids is not None: sync_prospect_projects(rid,project_ids,primary or '')
                    STATE.history(type_objet='prospect',objet_id=rid,objet_nom=s['nom_cible'],action='création depuis sourcing')
                STATE.client.table('sourcing').update({'importe':True,'prospect_id':rid}).eq('id',sid).execute(); return self.json_out({'ok':True,'prospect_id':rid})
            if path=='/api/import/inspect':
                n=int(self.headers.get('Content-Length','0')); raw=self.rfile.read(n); suffix=Path(self.headers.get('X-Filename','import.xlsx')).suffix or '.xlsx'; fd,tmp=tempfile.mkstemp(prefix='eurasante_import_',suffix=suffix); os.close(fd); Path(tmp).write_bytes(raw); token=os.path.basename(tmp); STATE.imports[token]=tmp; return self.json_out({'ok':True,'token':token,'sheets':inspect_excel(Path(tmp))})
            if path=='/api/import/apply':
                d=self.body_json(); tmp=STATE.imports.get(d['token']);
                if not tmp or not os.path.exists(tmp): raise ValueError('Import expiré. Re-sélectionne le fichier.')
                parsed=parse_excel(Path(tmp),d.get('mapping',{})); mode=d.get('mode','merge')
                if mode=='replace':
                    for t in ('prospect_projets','sourcing','prospects','projets'): STATE.client.table(t).delete().neq('id','00000000-0000-0000-0000-000000000000').execute()
                # Fusion non destructive : les cellules vides de l'Excel n'effacent jamais les donnees existantes.
                existing_p={norm(x['nom_sponsor']):x for x in STATE.fetch_all('prospects')}
                pc={'inserted':0,'updated':0}
                for row in parsed['prospects']:
                    k=norm(row['nom_sponsor']); ex=existing_p.get(k)
                    if ex:
                        upd=nonempty_update(row,ex)
                        oldraw=ex.get('donnees_source') or {}; newraw=row.get('donnees_source') or {}
                        upd['donnees_source']={**oldraw,**newraw}
                        STATE.client.table('prospects').update(upd).eq('id',ex['id']).execute(); pc['updated']+=1
                    else:
                        ins=dict(row)
                        ins['statut']=ins.get('statut') or '1. Prise de contact'
                        STATE.client.table('prospects').insert(ins).execute(); pc['inserted']+=1
                existing_j={norm(x['nom_projet']):x for x in STATE.fetch_all('projets')}
                jc={'inserted':0,'updated':0}
                for row in parsed['projects']:
                    k=norm(row['nom_projet']); ex=existing_j.get(k)
                    if ex:
                        upd=nonempty_update(row,ex)
                        oldraw=ex.get('donnees_source') or {}; newraw=row.get('donnees_source') or {}
                        upd['donnees_source']={**oldraw,**newraw}
                        STATE.client.table('projets').update(upd).eq('id',ex['id']).execute(); jc['updated']+=1
                    else: STATE.client.table('projets').insert(row).execute(); jc['inserted']+=1
                # Sourcing : mise a jour par couple source + nom au lieu de dupliquer a chaque reimport.
                existing_s={(x.get('source',''),norm(x.get('nom_cible',''))):x for x in STATE.fetch_all('sourcing')}
                sc={'inserted':0,'updated':0}
                for row in parsed['sourcing']:
                    key=(row.get('source',''),norm(row.get('nom_cible',''))); ex=existing_s.get(key)
                    if ex:
                        merged={**(ex.get('donnees') or {}),**(row.get('donnees') or {})}
                        STATE.client.table('sourcing').update({'donnees':merged}).eq('id',ex['id']).execute(); sc['updated']+=1
                    else:
                        STATE.client.table('sourcing').insert(row).execute(); sc['inserted']+=1
                # Alimente la table multiprojets depuis l'ancien champ Projet Cible après chaque import.
                backfill_legacy_project_links()
                try: os.remove(tmp)
                except Exception: pass
                STATE.imports.pop(d['token'],None)
                STATE.history(type_objet='import',objet_nom='Excel',action=f"import {mode}",nouvelle_valeur=f"prospects {pc}; projets {jc}; sourcing {sc}; stats {parsed.get('stats',{})}")
                return self.json_out({'ok':True,'prospects':pc,'projects':jc,'sourcing':sc,'stats':parsed.get('stats',{})})
            if path=='/api/relations/import-history':
                return self.json_out({'ok':True,**import_relation_history()})
            if path=='/api/relations/rdv/status':
                d=self.body_json(); rid=d.get('id'); statut=d.get('statut')
                if statut not in ('realise','planifie','annule','a_qualifier','ignore'): raise ValueError('Statut RDV invalide')
                STATE.client.table('rendez_vous').update({'statut':statut,'updated_at':datetime.now().isoformat()}).eq('id',rid).execute()
                return self.json_out({'ok':True})
            if path=='/api/relations/rdv/bulk-status':
                d=self.body_json(); ids=[str(x) for x in (d.get('ids') or []) if x]; statut=d.get('statut')
                if statut not in ('realise','planifie','annule','a_qualifier','ignore'): raise ValueError('Statut RDV invalide')
                if not ids: return self.json_out({'ok':True,'updated':0})
                updated=0
                # Mise à jour par lots : rapide pour plusieurs centaines de RDV.
                # Fallback ligne par ligne si PostgREST refuse un lot particulier.
                for pos in range(0,len(ids),100):
                    chunk=ids[pos:pos+100]
                    try:
                        resp=STATE.client.table('rendez_vous').update({'statut':statut,'updated_at':datetime.now().isoformat()}).in_('id',chunk).execute()
                        updated += len(resp.data or [])
                    except Exception:
                        for rid in chunk:
                            resp=STATE.client.table('rendez_vous').update({'statut':statut,'updated_at':datetime.now().isoformat()}).eq('id',rid).execute()
                            updated += len(resp.data or [])
                STATE.history(type_objet='relations',objet_nom='Rendez-vous',action='qualification en masse',nouvelle_valeur=f"{updated} RDV -> {statut}")
                return self.json_out({'ok':True,'updated':updated})
            if path=='/api/relations/organisation/status':
                d=self.body_json(); oid=d.get('id'); statut=d.get('statut')
                allowed=('a_qualifier','prospect','partenaire','ecosysteme','prestataire','institutionnel','ignorer')
                if statut not in allowed: raise ValueError('Type d organisation invalide')
                STATE.client.table('organisations').update({'statut_relation':statut,'updated_at':datetime.now().isoformat()}).eq('id',oid).execute()
                STATE.history(type_objet='relations',objet_id=oid,objet_nom='Organisation',action='qualification organisation',nouvelle_valeur=statut)
                return self.json_out({'ok':True})
            if path=='/api/relations/link-prospect':
                d=self.body_json(); oid=d.get('organisation_id'); pid=d.get('prospect_id') or None
                STATE.client.table('organisations').update({'prospect_id':pid,'statut_relation':'prospect','updated_at':datetime.now().isoformat()}).eq('id',oid).execute()
                return self.json_out({'ok':True})
            if path=='/api/relations/create-prospect':
                d=self.body_json(); oid=d.get('organisation_id')
                org=STATE.client.table('organisations').select('*').eq('id',oid).single().execute().data
                if not org: raise ValueError('Organisation introuvable')
                existing=STATE.client.table('prospects').select('*').ilike('nom_sponsor',org.get('nom','')).execute().data or []
                if existing: pid=existing[0]['id']
                else:
                    row=STATE.client.table('prospects').insert({'nom_sponsor':org.get('nom'),'statut':'1. Prise de contact','date_maj':date.today().isoformat(),'commentaire':'Créé depuis l’historique relationnel'}).execute().data[0]; pid=row['id']
                STATE.client.table('organisations').update({'prospect_id':pid,'statut_relation':'prospect','updated_at':datetime.now().isoformat()}).eq('id',oid).execute()
                return self.json_out({'ok':True,'prospect_id':pid})
            raise ValueError('Endpoint inconnu')
        except Exception as e: return self.fail(e)

if __name__=='__main__':
    os.chdir(ROOT)
    port=find_free_port()
    print('='*68); print('EURASANTE CRM - SUPABASE MULTI-UTILISATEUR'); print('='*68)
    print('Base centrale :',SUPABASE_URL); print('Interface :',f'http://127.0.0.1:{port}/app.html')
    print('Chaque poste se connecte avec son propre compte Supabase.')
    threading.Timer(0.9,lambda:webbrowser.open(f'http://127.0.0.1:{port}/app.html')).start()
    ThreadingHTTPServer(('127.0.0.1',port),Handler).serve_forever()
